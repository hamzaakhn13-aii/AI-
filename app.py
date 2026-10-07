"""AI Resume ATS Checker - Streamlit app powered by Google Gemini Flash."""

import io
import json
import re

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

MODEL_NAME = "gemini-2.5-flash"
MAX_CHARS = 30000  # safety cap on resume / job description length
MIN_CHARS = 150  # below this the file is probably a scanned image


# ----------------------------- File handling ----------------------------- #
def extract_text(uploaded_file) -> str:
    """Extract plain text from an uploaded PDF or DOCX file."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(io.BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password-protected.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        text = "\n".join(pages)
    elif name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.append(" | ".join(c.text for c in row.cells))
        text = "\n".join(parts)
    else:
        raise ValueError("Unsupported file type. Please upload a PDF or DOCX.")

    return re.sub(r"\n{3,}", "\n\n", text).strip()


# ------------------------------- AI logic -------------------------------- #
def build_prompt(resume_text: str, job_description: str) -> str:
    jd_block = (
        f"JOB DESCRIPTION:\n{job_description[:MAX_CHARS]}"
        if job_description.strip()
        else "JOB DESCRIPTION: Not provided. Evaluate against general ATS best "
        "practices and set keyword_match.match_percentage based on general "
        "industry keywords for the candidate's apparent field."
    )
    return f"""You are an expert ATS (Applicant Tracking System) analyst and resume coach.
Analyze the resume below and return ONLY a JSON object (no markdown) with this exact schema:

{{
  "ats_score": <integer 0-100>,
  "score_breakdown": {{
    "formatting": <integer 0-100>,
    "keywords": <integer 0-100>,
    "content_quality": <integer 0-100>,
    "readability": <integer 0-100>
  }},
  "summary": "<2-3 sentence overall assessment>",
  "strengths": ["<string>", ...],
  "weaknesses": ["<string>", ...],
  "keyword_match": {{
    "match_percentage": <integer 0-100>,
    "matched_keywords": ["<string>", ...],
    "missing_keywords": ["<string>", ...]
  }},
  "improvements": [
    {{"section": "<resume section>", "issue": "<what is wrong>", "suggestion": "<specific fix>", "priority": "High|Medium|Low"}}
  ],
  "rewritten_examples": [
    {{"original": "<weak bullet from resume>", "improved": "<stronger version>"}}
  ]
}}

Scoring guidance: be realistic and strict. Consider contact info, standard section headings,
quantified achievements, action verbs, keyword relevance, length, consistency, and spelling.
Give 4-8 improvements ordered by priority and 2-4 rewritten examples using only facts
present in the resume (never invent numbers or employers).

{jd_block}

RESUME:
{resume_text[:MAX_CHARS]}
"""


def parse_json_response(raw: str) -> dict:
    """Parse model output into a dict, tolerating markdown fences / extra text."""
    raw = (raw or "").strip()
    raw = re.sub(r"^```(?:json)?\s*|\s*```$", "", raw, flags=re.IGNORECASE)
    try:
        return json.loads(raw)
    except json.JSONDecodeError:
        match = re.search(r"\{.*\}", raw, re.DOTALL)
        if match:
            return json.loads(match.group(0))
        raise ValueError("The AI returned an unreadable response. Please try again.")


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def normalize_result(data: dict) -> dict:
    """Fill defaults and clamp values so the UI never crashes on odd output."""
    breakdown = data.get("score_breakdown") or {}
    km = data.get("keyword_match") or {}
    return {
        "ats_score": _clamp(data.get("ats_score")),
        "score_breakdown": {
            k: _clamp(breakdown.get(k))
            for k in ("formatting", "keywords", "content_quality", "readability")
        },
        "summary": str(data.get("summary", "")),
        "strengths": list(data.get("strengths") or []),
        "weaknesses": list(data.get("weaknesses") or []),
        "keyword_match": {
            "match_percentage": _clamp(km.get("match_percentage")),
            "matched_keywords": list(km.get("matched_keywords") or []),
            "missing_keywords": list(km.get("missing_keywords") or []),
        },
        "improvements": [i for i in (data.get("improvements") or []) if isinstance(i, dict)],
        "rewritten_examples": [
            i for i in (data.get("rewritten_examples") or []) if isinstance(i, dict)
        ],
    }


def analyze_resume(api_key: str, resume_text: str, job_description: str) -> dict:
    client = genai.Client(api_key=api_key)
    response = client.models.generate_content(
        model=MODEL_NAME,
        contents=build_prompt(resume_text, job_description),
        config=types.GenerateContentConfig(
            temperature=0.2,
            response_mime_type="application/json",
        ),
    )
    return normalize_result(parse_json_response(response.text))


# --------------------------------- UI ------------------------------------ #
def score_color(score: int) -> str:
    return "#16a34a" if score >= 75 else "#f59e0b" if score >= 50 else "#dc2626"


def render_results(result: dict) -> None:
    score = result["ats_score"]
    st.markdown(
        f"""
        <div style="text-align:center;padding:20px;border-radius:16px;
                    border:2px solid {score_color(score)};margin-bottom:16px;">
            <div style="font-size:64px;font-weight:800;color:{score_color(score)};">{score}/100</div>
            <div style="font-size:18px;">Overall ATS Score</div>
        </div>
        """,
        unsafe_allow_html=True,
    )
    st.write(result["summary"])

    cols = st.columns(4)
    labels = {
        "formatting": "Formatting",
        "keywords": "Keywords",
        "content_quality": "Content",
        "readability": "Readability",
    }
    for col, (key, label) in zip(cols, labels.items()):
        val = result["score_breakdown"][key]
        col.metric(label, f"{val}/100")
        col.progress(val / 100)

    c1, c2 = st.columns(2)
    with c1:
        st.subheader("✅ Strengths")
        for s in result["strengths"] or ["None identified."]:
            st.markdown(f"- {s}")
    with c2:
        st.subheader("⚠️ Weaknesses")
        for w in result["weaknesses"] or ["None identified."]:
            st.markdown(f"- {w}")

    km = result["keyword_match"]
    st.subheader(f"🔑 Keyword Match: {km['match_percentage']}%")
    k1, k2 = st.columns(2)
    with k1:
        st.markdown("**Found**")
        st.write(", ".join(km["matched_keywords"]) or "—")
    with k2:
        st.markdown("**Missing**")
        st.write(", ".join(km["missing_keywords"]) or "—")

    st.subheader("🛠️ Recommended Improvements")
    icons = {"high": "🔴", "medium": "🟠", "low": "🟢"}
    for imp in result["improvements"]:
        prio = str(imp.get("priority", "Medium"))
        icon = icons.get(prio.lower(), "🟠")
        with st.expander(f"{icon} [{prio}] {imp.get('section', 'General')}"):
            st.markdown(f"**Issue:** {imp.get('issue', '')}")
            st.markdown(f"**Fix:** {imp.get('suggestion', '')}")

    if result["rewritten_examples"]:
        st.subheader("✍️ Example Rewrites")
        for ex in result["rewritten_examples"]:
            st.markdown(f"❌ *{ex.get('original', '')}*")
            st.markdown(f"✅ **{ex.get('improved', '')}**")
            st.divider()

    st.download_button(
        "⬇️ Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def get_api_key() -> str:
    """Prefer Streamlit secrets (deployment); fall back to sidebar input."""
    try:
        if "GEMINI_API_KEY" in st.secrets:
            return st.secrets["GEMINI_API_KEY"]
    except Exception:
        pass  # no secrets file locally
    return st.sidebar.text_input(
        "Gemini API Key",
        type="password",
        help="Get a free key at https://aistudio.google.com/apikey",
    )


def main() -> None:
    st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="wide")
    st.title("📄 AI Resume ATS Checker")
    st.caption("Upload your resume to get an ATS score and actionable improvements.")

    api_key = get_api_key()

    left, right = st.columns(2)
    with left:
        uploaded = st.file_uploader("Upload resume (PDF or DOCX)", type=["pdf", "docx"])
    with right:
        jd = st.text_area(
            "Job description (optional, improves keyword analysis)", height=180
        )

    if st.button("Analyze Resume", type="primary", use_container_width=True):
        if not api_key:
            st.error("Please provide your Gemini API key in the sidebar.")
            return
        if uploaded is None:
            st.error("Please upload a resume first.")
            return
        try:
            text = extract_text(uploaded)
        except Exception as e:
            st.error(f"Could not read the file: {e}")
            return
        if len(text) < MIN_CHARS:
            st.error(
                "Very little text was found. The file may be a scanned image, "
                "which ATS systems also cannot read. Export a text-based PDF or DOCX."
            )
            return
        try:
            with st.spinner("Analyzing your resume..."):
                st.session_state["result"] = analyze_resume(api_key, text, jd)
        except Exception as e:
            st.error(f"Analysis failed: {e}")
            return

    if "result" in st.session_state:
        st.divider()
        render_results(st.session_state["result"])


if __name__ == "__main__":
    main()
